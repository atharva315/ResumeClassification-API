from __future__ import annotations
import io
import json
import os
import re
import time
import unicodedata
from contextlib import asynccontextmanager
from pathlib import Path
from typing import List, Optional, Tuple
import torch
from fastapi import FastAPI, File, HTTPException, UploadFile
from pydantic import BaseModel, Field
from transformers import AutoModelForSequenceClassification, AutoTokenizer
import pymupdf  # Use the current PyMuPDF module name (avoids deprecated fitz import)
import pdfplumber
import pytesseract
from docx import Document
from PIL import Image, ImageOps, UnidentifiedImageError
from prometheus_client import (
    CONTENT_TYPE_LATEST,
    Counter,
    Gauge,
    Histogram,
    generate_latest,
)
from starlette.requests import Request
from starlette.responses import Response
# ============================================================
# Configuration
# ============================================================
BASE_DIR = Path(__file__).resolve().parent
MODEL_DIR = BASE_DIR / "model"
MAX_LENGTH = int(os.getenv("MODEL_MAX_LENGTH", "2048"))
# Provisional softmax-score gate. Validate with validation data before production.
MIN_PREDICTION_SCORE = float(os.getenv("MIN_PREDICTION_SCORE", "0.25"))
MIN_RESUME_WORDS = int(os.getenv("MIN_RESUME_WORDS", "30"))
MODEL_NAME = "ModernBERT Base — Resume Classification"
DEVICE = torch.device("cuda" if torch.cuda.is_available() else "cpu")
MAX_UPLOAD_BYTES = 10 * 1024 * 1024  # 10 MB
MAX_PDF_PAGES = 10
MAX_IMAGE_PIXELS = 35_000_000
MAX_TEXT_CHARS = 200_000
ALLOWED_EXTENSIONS = {
    ".pdf", ".txt", ".docx",
    ".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff",
}
# Optional override, e.g. TESSERACT_CMD=C:\Program Files\Tesseract-OCR\tesseract.exe
TESSERACT_CMD = os.getenv("TESSERACT_CMD")
if TESSERACT_CMD:
    pytesseract.pytesseract.tesseract_cmd = TESSERACT_CMD
else:
    default_tesseract = Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe")
    if default_tesseract.exists():
        pytesseract.pytesseract.tesseract_cmd = str(default_tesseract)
# PSM 4 performed better than PSM 6 for the user's supplied single-column resume.
# Override through TESSERACT_PSM if another layout needs testing.
TESSERACT_PSM = int(os.getenv("TESSERACT_PSM", "4"))
if not 0.0 <= MIN_PREDICTION_SCORE <= 1.0:
    raise ValueError("MIN_PREDICTION_SCORE must be between 0 and 1.")
if MAX_LENGTH < 1:
    raise ValueError("MODEL_MAX_LENGTH must be a positive integer.")
if MIN_RESUME_WORDS < 1:
    raise ValueError("MIN_RESUME_WORDS must be a positive integer.")
# Prometheus metrics. Labels use fixed, low-cardinality values only; no resume
# text, filenames, user data, or predicted category is stored in metrics.
HTTP_REQUESTS = Counter(
    "resume_api_http_requests",
    "HTTP requests served by the Resume Classification API.",
    ["method", "endpoint", "status_code"],
)
HTTP_REQUEST_DURATION = Histogram(
    "resume_api_http_request_duration_seconds",
    "End-to-end HTTP request duration, including file extraction when applicable.",
    ["method", "endpoint"],
    buckets=(0.01, 0.025, 0.05, 0.1, 0.25, 0.5, 0.75, 1.0, 1.5, 2.5, 5.0, 10.0, 30.0),
)
PREDICTION_RESULTS = Counter(
    "resume_api_prediction_results",
    "Prediction outcomes produced by the API.",
    ["result"],
)
PREDICTION_STAGE_DURATION = Histogram(
    "resume_api_stage_duration_seconds",
    "Duration of prediction pipeline stages in seconds.",
    ["stage"],
    buckets=(0.005, 0.01, 0.025, 0.05, 0.1, 0.25, 0.5, 0.75, 1.0, 1.5, 2.5, 5.0, 10.0, 30.0),
)
MODEL_READY = Gauge(
    "resume_api_model_ready",
    "1 when the model is loaded and ready to serve predictions, otherwise 0.",
)
MODEL_READY.set(0)
if torch.cuda.is_available():
    torch.backends.cuda.matmul.allow_tf32 = True
    torch.backends.cudnn.allow_tf32 = True
tokenizer = None
model = None
id2label = None
label2id = None
# ============================================================
# Model startup/lifecycle
# ============================================================
@asynccontextmanager
async def lifespan(app: FastAPI):
    global tokenizer, model, id2label, label2id
    MODEL_READY.set(0)
    mapping_file = MODEL_DIR / "label_mapping.json"
    if not mapping_file.exists():
        raise RuntimeError(f"Label mapping is missing: {mapping_file}")
    with mapping_file.open("r", encoding="utf-8") as file:
        mapping = json.load(file)
    label2id = {
        str(label): int(index)
        for label, index in mapping["label2id"].items()
    }
    id2label = {index: label for label, index in label2id.items()}
    if len(label2id) != 43:
        raise RuntimeError(f"Expected 43 categories, found {len(label2id)}.")
    if sorted(label2id.values()) != list(range(43)):
        raise RuntimeError("Label IDs must be consecutive integers from 0 to 42.")
    tokenizer = AutoTokenizer.from_pretrained(
        str(MODEL_DIR), local_files_only=True
    )
    tokenizer.padding_side = "right"
    model = AutoModelForSequenceClassification.from_pretrained(
        str(MODEL_DIR), local_files_only=True
    ).to(DEVICE)
    model.eval()
    if model.config.num_labels != len(label2id):
        raise RuntimeError("Model output count does not match the label mapping.")
    model_max_length = getattr(model.config, "max_position_embeddings", 8192)
    if MAX_LENGTH > model_max_length:
        raise RuntimeError(
            f"MODEL_MAX_LENGTH ({MAX_LENGTH}) exceeds model capacity ({model_max_length})."
        )
    MODEL_READY.set(1)
    print(
        f"Loaded {MODEL_NAME} on {DEVICE}; max tokens={MAX_LENGTH}; "
        f"minimum score={MIN_PREDICTION_SCORE:.0%}; OCR PSM={TESSERACT_PSM}"
    )
    yield
    MODEL_READY.set(0)
    model = None
    tokenizer = None
    if torch.cuda.is_available():
        torch.cuda.empty_cache()
app = FastAPI(
    title="Resume Classification API",
    description=(
        "ModernBERT resume classification with a minimum-score gate, "
        "four ranked alternatives."
    ),
    version="1.10.0",
    lifespan=lifespan,
)
@app.middleware("http")
async def collect_prometheus_http_metrics(request: Request, call_next):
    """Record low-cardinality request count and latency metrics."""
    # Do not count Prometheus scraping itself, otherwise dashboard traffic
    # would inflate application request volume.
    if request.url.path == "/metrics":
        return await call_next(request)
    started = time.perf_counter()
    status_code = 500
    try:
        response = await call_next(request)
        status_code = response.status_code
        return response
    finally:
        route = request.scope.get("route")
        endpoint = getattr(route, "path", "/unmatched") or "/unmatched"
        HTTP_REQUESTS.labels(
            method=request.method,
            endpoint=endpoint,
            status_code=str(status_code),
        ).inc()
        HTTP_REQUEST_DURATION.labels(
            method=request.method,
            endpoint=endpoint,
        ).observe(time.perf_counter() - started)
# ============================================================
# API schemas
# ============================================================
class PredictRequest(BaseModel):
    prepared_text: str = Field(
        ...,
        min_length=1,
        max_length=MAX_TEXT_CHARS,
        description=(
            "Resume text. The API applies conservative Unicode and whitespace "
            "normalization while preserving meaningful words, numbers, punctuation, "
            "and line breaks."
        ),
    )
class ClassPrediction(BaseModel):
    category: str
    probability: float
class PredictResponse(BaseModel):
    predicted_category: Optional[str] = None
    prediction_message: str
    other_categories: List[ClassPrediction] = Field(default_factory=list)
    alternatives_message: str
    token_count_before_truncation: int
    input_limit_tokens: int
    was_truncated: bool
    inference_ms: float
    prediction_ms: float
    model: str
# ============================================================
# Resume suitability gate
# ============================================================
WORD_RE = re.compile(r"\b[A-Za-z][A-Za-z0-9+#./-]*\b")
RESUME_SIGNAL_RE = re.compile(
    r"\b(?:education|skills?|experience|employment|projects?|certifications?|"
    r"professional\s+(?:summary|profile)|work\s+history|internship|degree|"
    r"university|college|bachelor|master|developer|engineer|analyst|manager|"
    r"accountant|designer|architect|consultant|programmer|teacher|researcher|"
    r"responsibilities|achievements|career|qualification|qualifications|"
    r"built|developed|implemented|designed|managed|led|deployed|intern)\b",
    re.IGNORECASE,
)
def resume_input_check(text: str) -> Tuple[bool, str, int]:
    """Lightweight guard against short, repetitive, or obviously unrelated input."""
    words = WORD_RE.findall(text)
    word_count = len(words)
    if word_count < MIN_RESUME_WORDS:
        return (
            False,
            f"Please provide a fuller resume. At least {MIN_RESUME_WORDS} words are required.",
            word_count,
        )
    if not RESUME_SIGNAL_RE.search(text):
        return (
            False,
            "This text does not look like resume content. Include skills, education, work experience, projects, or qualifications.",
            word_count,
        )
    unique_ratio = len({word.lower() for word in words}) / max(word_count, 1)
    if unique_ratio < 0.15:
        return (
            False,
            "This input appears too repetitive to classify as a resume. Please provide genuine resume content.",
            word_count,
        )
    return True, "", word_count
# ============================================================
# Input normalization and file extraction
# ============================================================
def normalize_resume_text(text: str) -> str:
    """Normalize extraction artifacts without aggressively changing content."""
    if not text:
        return ""
    text = unicodedata.normalize("NFKC", str(text))
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # Remove control/format characters but preserve line breaks and tabs.
    text = "".join(
        char for char in text
        if char in "\n\t" or unicodedata.category(char)[0] != "C"
    )
    # Collapse horizontal whitespace while retaining section boundaries.
    lines = [re.sub(r"[ \t\f\v]+", " ", line).strip() for line in text.split("\n")]
    text = "\n".join(lines)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) > MAX_TEXT_CHARS:
        raise HTTPException(
            status_code=413,
            detail=f"Extracted resume text exceeds {MAX_TEXT_CHARS:,} characters.",
        )
    return text
def run_ocr(image: Image.Image) -> str:
    """OCR a PIL image using PSM 4, validated for this resume's single-column layout."""
    try:
        image = ImageOps.exif_transpose(image).convert("RGB")
        return pytesseract.image_to_string(
            image,
            lang="eng",
            config=f"--oem 3 --psm {TESSERACT_PSM}",
        )
    except pytesseract.pytesseract.TesseractNotFoundError as exc:
        raise HTTPException(
            status_code=503,
            detail=(
                "OCR engine is not installed or not on PATH. Install Tesseract OCR "
                "and restart the terminal/API, or configure TESSERACT_CMD."
            ),
        ) from exc
def _ocr_pdf_pages(pdf_bytes: bytes, page_indices: List[int]) -> dict[int, str]:
    """OCR scanned PDF pages while opening the PDF only once."""
    results: dict[int, str] = {}
    try:
        with pymupdf.open(stream=pdf_bytes, filetype="pdf") as pdf_doc:
            for page_index in page_indices:
                page = pdf_doc.load_page(page_index)
                pixmap = page.get_pixmap(
                    dpi=160,
                    colorspace=pymupdf.csRGB,
                    alpha=False,
                )
                image_bytes = pixmap.tobytes("png")
                with Image.open(io.BytesIO(image_bytes)) as image:
                    image.load()
                    if image.width * image.height > MAX_IMAGE_PIXELS:
                        raise HTTPException(
                            status_code=413,
                            detail="A scanned PDF page is too large to process safely.",
                        )
                    results[page_index] = run_ocr(image)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=422,
            detail=f"Could not OCR scanned PDF page ({type(exc).__name__}).",
        ) from exc
    return results
def extract_resume_text(filename: str, data: bytes) -> str:
    """Extract text from PDF, TXT, DOCX, or supported image formats."""
    suffix = Path(filename or "").suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=415,
            detail=(
                "Unsupported file type. Upload PDF, TXT, DOCX, PNG, JPG, JPEG, "
                "WEBP, BMP, TIF, or TIFF."
            ),
        )
    if not data:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")
    try:
        if suffix == ".txt":
            extracted = data.decode("utf-8-sig", errors="replace")
        elif suffix == ".docx":
            document = Document(io.BytesIO(data))
            parts = [
                paragraph.text for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]
            # Many resume templates use tables for columns/sections.
            for table in document.tables:
                for row in table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        parts.append(row_text)
            extracted = "\n".join(parts)
        elif suffix == ".pdf":
            if b"%PDF-" not in data[:1024]:
                raise HTTPException(
                    status_code=415,
                    detail="The uploaded file does not appear to be a valid PDF.",
                )
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                if len(pdf.pages) > MAX_PDF_PAGES:
                    raise HTTPException(
                        status_code=413,
                        detail=f"PDFs are limited to {MAX_PDF_PAGES} pages for this demo.",
                    )
                page_texts = [page.extract_text() or "" for page in pdf.pages]
            # OCR only pages with no usable text layer.
            pages_needing_ocr = [
                index for index, page_text in enumerate(page_texts)
                if len(page_text.strip()) < 25
            ]
            if pages_needing_ocr:
                ocr_results = _ocr_pdf_pages(data, pages_needing_ocr)
                for page_index, ocr_text in ocr_results.items():
                    if len(ocr_text.strip()) > len(page_texts[page_index].strip()):
                        page_texts[page_index] = ocr_text
            extracted = "\n\n".join(page_texts)
        else:
            # Verify the image from its actual contents, not its extension alone.
            with Image.open(io.BytesIO(data)) as image:
                image.load()
                if image.width * image.height > MAX_IMAGE_PIXELS:
                    raise HTTPException(
                        status_code=413,
                        detail=f"Images must be no larger than {MAX_IMAGE_PIXELS:,} pixels.",
                    )
                extracted = run_ocr(image)
    except HTTPException:
        raise
    except pytesseract.pytesseract.TesseractNotFoundError as exc:
        raise HTTPException(
            status_code=503,
            detail="OCR is unavailable because Tesseract is not installed or not on PATH.",
        ) from exc
    except UnidentifiedImageError as exc:
        raise HTTPException(
            status_code=415,
            detail="The uploaded image is invalid or its format is unsupported.",
        ) from exc
    except Exception as exc:
        raise HTTPException(
            status_code=422,
            detail=f"Could not read the uploaded file ({type(exc).__name__}).",
        ) from exc
    extracted = normalize_resume_text(extracted)
    if not extracted:
        raise HTTPException(
            status_code=422,
            detail=(
                "No readable text was found. For scanned files/images, use a clearer image "
                "and ensure Tesseract OCR is installed."
            ),
        )
    return extracted
# ============================================================
# Batch model inference
# ============================================================
def get_probabilities_batch(texts: List[str]) -> List[torch.Tensor]:
    """Tokenize and evaluate one or more resume texts."""
    if not texts:
        return []
    encoded = tokenizer(
        texts,
        truncation=True,
        max_length=MAX_LENGTH,
        padding=True,
        return_attention_mask=True,
        return_tensors="pt",
    )
    inputs = {key: value.to(DEVICE) for key, value in encoded.items()}
    batch_size = min(len(texts), 4 if DEVICE.type == "cuda" else 2)
    while True:
        outputs = []
        try:
            for start in range(0, len(texts), batch_size):
                batch_inputs = {
                    key: value[start:start + batch_size]
                    for key, value in inputs.items()
                }
                with torch.inference_mode():
                    if DEVICE.type == "cuda":
                        with torch.autocast(device_type="cuda", dtype=torch.float16):
                            logits = model(**batch_inputs).logits
                    else:
                        logits = model(**batch_inputs).logits
                    probabilities = torch.softmax(logits.float(), dim=-1).cpu()
                outputs.extend(probabilities.unbind(0))
            return outputs
        except torch.cuda.OutOfMemoryError:
            if DEVICE.type != "cuda" or batch_size <= 1:
                raise
            outputs.clear()
            torch.cuda.empty_cache()
            batch_size = max(1, batch_size // 2)
def get_probabilities(text: str) -> torch.Tensor:
    return get_probabilities_batch([text])[0]
def count_tokens(text: str) -> int:
    return len(tokenizer(text, truncation=False, add_special_tokens=True)["input_ids"])
# ============================================================
# Consistent rejection responses
# ============================================================
def rejection_response(
    message: str,
    token_count: int,
    total_start: float,
    prediction_ms: float = 0.0,
) -> PredictResponse:
    elapsed = (time.perf_counter() - total_start) * 1000.0
    PREDICTION_STAGE_DURATION.labels(stage="rejected_pipeline").observe(
        elapsed / 1000.0
    )
    return PredictResponse(
        predicted_category=None,
        prediction_message=message,
        other_categories=[],
        alternatives_message="Alternative categories are hidden because the input was not accepted.",
        token_count_before_truncation=token_count,
        input_limit_tokens=MAX_LENGTH,
        was_truncated=token_count > MAX_LENGTH,
        inference_ms=round(elapsed, 2),
        prediction_ms=round(prediction_ms, 2),
        model=MODEL_NAME,
    )
# ============================================================
# Endpoints
# ============================================================
@app.get("/health")
def health():
    return {
        "status": "ok" if model is not None else "loading",
        "model": MODEL_NAME,
        "device": str(DEVICE),
        "max_input_tokens": MAX_LENGTH,
        "classes": len(label2id) if label2id else 0,
        "minimum_prediction_score": MIN_PREDICTION_SCORE,
        "tesseract_psm": TESSERACT_PSM,
    }
@app.get("/")
def root():
    return {
        "message": "Resume Classification API",
        "docs": "/docs",
        "health": "/health",
    }
@app.get("/metrics", include_in_schema=False)
def prometheus_metrics():
    """Expose Prometheus metrics. Restrict this endpoint to trusted monitoring networks in production."""
    return Response(generate_latest(), media_type=CONTENT_TYPE_LATEST)
@app.post("/predict", response_model=PredictResponse)
def predict(request: PredictRequest):
    if model is None or tokenizer is None:
        raise HTTPException(status_code=503, detail="Model is not ready.")
    text = normalize_resume_text(request.prepared_text)
    total_start = time.perf_counter()
    if not text:
        raise HTTPException(status_code=422, detail="prepared_text must not be blank.")
    try:
        original_token_count = count_tokens(text)
        looks_like_resume, rejection_message, _ = resume_input_check(text)
        if not looks_like_resume:
            PREDICTION_RESULTS.labels(result="rejected_input").inc()
            return rejection_response(
                rejection_message,
                token_count=original_token_count,
                total_start=total_start,
            )
        was_truncated = original_token_count > MAX_LENGTH
        prediction_start = time.perf_counter()
        probabilities = get_probabilities(text)
        prediction_ms = (time.perf_counter() - prediction_start) * 1000.0
        PREDICTION_STAGE_DURATION.labels(stage="prediction").observe(
            prediction_ms / 1000.0
        )
        all_probs, all_indices = torch.sort(probabilities, descending=True)
        predicted_index = int(all_indices[0].item())
        top_score = float(all_probs[0].item())
        if top_score < MIN_PREDICTION_SCORE:
            PREDICTION_RESULTS.labels(result="rejected_score").inc()
            return rejection_response(
                (
                    "We couldn't confidently classify this resume. Please review the "
                    "text and ensure it includes relevant skills, experience, education, or projects."
                ),
                token_count=original_token_count,
                total_start=total_start,
                prediction_ms=prediction_ms,
            )
        predicted_category = id2label[predicted_index]
        alternatives = [
            ClassPrediction(
                category=id2label[int(index)],
                probability=round(float(probability), 6),
            )
            for probability, index in zip(
                all_probs[1:5].tolist(),
                all_indices[1:5].tolist(),
            )
        ]
        PREDICTION_RESULTS.labels(result="accepted").inc()
        total_inference_ms = (time.perf_counter() - total_start) * 1000.0
        PREDICTION_STAGE_DURATION.labels(stage="full_inference").observe(
            total_inference_ms / 1000.0
        )
        return PredictResponse(
            predicted_category=predicted_category,
            prediction_message="The prediction passed the minimum score threshold.",
            other_categories=alternatives,
            alternatives_message=(
                "The four highest-scoring alternative categories are listed "
                "in descending score order."
            ),
            token_count_before_truncation=original_token_count,
            input_limit_tokens=MAX_LENGTH,
            was_truncated=was_truncated,
            inference_ms=round(total_inference_ms, 2),
            prediction_ms=round(prediction_ms, 2),
            model=MODEL_NAME,
        )
    except HTTPException:
        raise
    except Exception as exc:
        PREDICTION_RESULTS.labels(result="error").inc()
        print(f"Prediction failed: {type(exc).__name__}: {exc}")
        raise HTTPException(
            status_code=500,
            detail=f"Inference failed: {type(exc).__name__}",
        ) from exc
@app.post("/predict-file", response_model=PredictResponse)
def predict_file(file: UploadFile = File(...)):
    """Upload PDF/TXT/DOCX/image and classify its extracted text."""
    filename = file.filename or "resume"
    try:
        data = file.file.read(MAX_UPLOAD_BYTES + 1)
        if len(data) > MAX_UPLOAD_BYTES:
            raise HTTPException(
                status_code=413,
                detail=f"Upload is too large. Maximum file size is {MAX_UPLOAD_BYTES // (1024 * 1024)} MB.",
            )
        extracted_text = extract_resume_text(filename, data)
        # Reuse the identical text normalization, suitability, and prediction path.
        return predict(PredictRequest(prepared_text=extracted_text))
    finally:
        try:
            file.file.close()
        except Exception:
            pass