"""Automated tests for the current Resume Classification FastAPI service.

The tests mock model loading/inference, so they run without the large model
checkpoint. Parser tests use synthetic content created in memory.
Run from services/api: python -m pytest -q tests
"""

import io

import pytest
import torch
from docx import Document
from fastapi.testclient import TestClient
import pymupdf
from PIL import Image

import app as api


LABELS = [
    "Accountant", "Advocate", "Agriculture", "Apparel", "Architecture",
    "Arts", "Automobile", "Aviation", "BPO", "Banking", "Blockchain",
    "Building and Construction", "Business Analyst", "Civil Engineer",
    "Consultant", "Data Science", "Database", "Designing", "DevOps",
    "Digital Media", "DotNet Developer", "ETL Developer", "Education",
    "Electrical Engineering", "Finance", "Food and Beverages",
    "Health and Fitness", "Human Resources", "Information Technology",
    "Java Developer", "Management", "Mechanical Engineer",
    "Network Security Engineer", "Operations Manager", "PMO",
    "Public Relations", "Python Developer", "React Developer",
    "SAP Developer", "SQL Developer", "Sales", "Testing", "Web Designing",
]

VALID_RESUME = """Professional Summary
Python developer with experience building backend applications.
Skills: Python, FastAPI, Django, REST API testing, debugging, Git, SQL.
Education: Bachelor's degree in engineering.
Experience: Developed services, implemented API integrations, deployed applications,
wrote unit tests, maintained code, and collaborated with teams on projects.
"""


def strong_probabilities():
    """Return deterministic scores with Python Developer as the top class."""
    values = torch.full((43,), 0.14 / 38, dtype=torch.float32)
    values[36] = 0.40  # Python Developer
    values[15] = 0.20  # Data Science
    values[20] = 0.12  # DotNet Developer
    values[29] = 0.08  # Java Developer
    values[21] = 0.06  # ETL Developer
    return values


@pytest.fixture
def client(monkeypatch):
    """Set fake loaded-model state without loading real model weights."""
    labels = {i: label for i, label in enumerate(LABELS)}
    monkeypatch.setattr(api, "model", object())
    monkeypatch.setattr(api, "tokenizer", object())
    monkeypatch.setattr(api, "id2label", labels)
    monkeypatch.setattr(api, "label2id", {value: key for key, value in labels.items()})
    monkeypatch.setattr(api, "count_tokens", lambda text: len(text.split()) + 2)
    monkeypatch.setattr(api, "get_probabilities", lambda text: strong_probabilities())

    test_client = TestClient(api.app)
    yield test_client
    test_client.close()


def test_normalize_resume_text_masks_contact_details_like_training():
    source = (
        "Contact: jane.doe@example.com\\n"
        "Website: https://example.com/profile\\n"
        "LinkedIn: linkedincom/in/janedoe\\n"
        "Phone: +1 (555) 123-4567\\n"
        "Skills: Python, C++, C#, and SQL"
    )

    normalized = api.normalize_resume_text(source)

    assert "[EMAIL]" in normalized
    assert "[URL]" in normalized
    assert "[PHONE]" in normalized
    assert "jane.doe@example.com" not in normalized
    assert "https://example.com/profile" not in normalized
    assert "+1 (555) 123-4567" not in normalized
    assert "\\n" not in normalized
    assert "C++" in normalized
    assert "C#" in normalized
    assert "SQL" in normalized


def test_normalize_resume_text_masks_flattened_contact_patterns():
    normalized = api.normalize_resume_text(
        "Email johnsmithgmailcom URL www.example.org Phone 9876543210"
    )

    assert "[EMAIL]" in normalized
    assert "[URL]" in normalized
    assert "[PHONE]" in normalized
    assert "johnsmithgmailcom" not in normalized
    assert "www.example.org" not in normalized
    assert "9876543210" not in normalized


def test_root_endpoint(client):
    response = client.get("/")
    assert response.status_code == 200
    assert response.json()["docs"] == "/docs"
    assert response.json()["health"] == "/health"


def test_health_endpoint_when_model_is_ready(client):
    response = client.get("/health")
    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "ok"
    assert body["classes"] == 43
    assert body["max_input_tokens"] == api.MAX_LENGTH


def test_predict_returns_primary_category_and_four_ranked_alternatives(client):
    response = client.post("/predict", json={"prepared_text": VALID_RESUME})
    assert response.status_code == 200
    body = response.json()

    assert body["predicted_category"] == "Python Developer"
    assert len(body["other_categories"]) == 4
    assert [item["category"] for item in body["other_categories"]] == [
        "Data Science", "DotNet Developer", "Java Developer", "ETL Developer"
    ]
    scores = [item["probability"] for item in body["other_categories"]]
    assert scores == sorted(scores, reverse=True)
    assert body["prediction_ms"] >= 0
    assert body["inference_ms"] >= 0
    assert body["input_limit_tokens"] == api.MAX_LENGTH


def test_low_score_hides_prediction_and_alternatives(client, monkeypatch):
    monkeypatch.setattr(
        api,
        "get_probabilities",
        lambda text: torch.full((43,), 1.0 / 43, dtype=torch.float32),
    )
    response = client.post("/predict", json={"prepared_text": VALID_RESUME})
    assert response.status_code == 200
    body = response.json()
    assert body["predicted_category"] is None
    assert body["other_categories"] == []
    assert "hidden" in body["alternatives_message"].lower()


def test_irrelevant_repetitive_text_is_rejected_before_inference(client, monkeypatch):
    calls = []

    def unexpected_inference(text):
        calls.append(text)
        raise AssertionError("Model inference should not run for unrelated input")

    monkeypatch.setattr(api, "get_probabilities", unexpected_inference)
    response = client.post(
        "/predict",
        json={"prepared_text": ("banana " * 80).strip()},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["predicted_category"] is None
    assert body["other_categories"] == []
    assert calls == []


def test_missing_text_is_rejected_by_validation(client):
    response = client.post("/predict", json={})
    assert response.status_code == 422


def test_multiline_raw_json_with_unescaped_newline_is_rejected(client):
    response = client.post(
        "/predict",
        content='{"prepared_text": "line one\nline two"}',
        headers={"Content-Type": "application/json"},
    )
    assert response.status_code == 422


def test_txt_file_upload_uses_prediction_pipeline(client, monkeypatch):
    monkeypatch.setattr(api, "extract_resume_text", lambda filename, data: VALID_RESUME)
    response = client.post(
        "/predict-file",
        files={"file": ("resume.txt", b"synthetic test content", "text/plain")},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["predicted_category"] == "Python Developer"
    assert len(body["other_categories"]) == 4


def test_unsupported_upload_type_returns_415(client):
    response = client.post(
        "/predict-file",
        files={"file": ("resume.exe", b"not a resume", "application/octet-stream")},
    )
    assert response.status_code == 415


def test_txt_extraction():
    text = api.extract_resume_text(
        "resume.txt", b"Skills: Python\r\nExperience: backend development"
    )
    assert "Skills: Python" in text
    assert "Experience: backend development" in text


def test_docx_extraction_includes_paragraphs_and_table():
    document = Document()
    document.add_paragraph("Experience: Python backend developer")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Skills"
    table.cell(0, 1).text = "Python FastAPI"

    buffer = io.BytesIO()
    document.save(buffer)
    extracted = api.extract_resume_text("resume.docx", buffer.getvalue())

    assert "Experience: Python backend developer" in extracted
    assert "Python FastAPI" in extracted


def test_pdf_text_extraction():
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Experience Python developer skills FastAPI")
    pdf_bytes = document.tobytes()
    document.close()

    extracted = api.extract_resume_text("resume.pdf", pdf_bytes)
    assert "Experience" in extracted
    assert "Python" in extracted


def test_image_extraction_routes_through_ocr(monkeypatch):
    monkeypatch.setattr(api, "run_ocr", lambda image: "Skills: Python and FastAPI")

    image_buffer = io.BytesIO()
    Image.new("RGB", (200, 100), color="white").save(image_buffer, format="PNG")
    extracted = api.extract_resume_text("resume.png", image_buffer.getvalue())

    assert "Skills: Python and FastAPI" in extracted


def test_run_ocr_uses_configured_layout_mode(monkeypatch):
    captured = {}

    def fake_ocr(image, lang, config):
        captured["lang"] = lang
        captured["config"] = config
        return "Recognized resume text"

    monkeypatch.setattr(api.pytesseract, "image_to_string", fake_ocr)
    result = api.run_ocr(Image.new("RGB", (100, 80), color="white"))

    assert result == "Recognized resume text"
    assert captured["lang"] == "eng"
    assert captured["config"] == f"--oem 3 --psm {api.TESSERACT_PSM}"


def test_prometheus_metrics_endpoint(client):
    assert client.get("/").status_code == 200
    response = client.get("/metrics")
    assert response.status_code == 200
    assert "resume_api_http_requests_total" in response.text
    assert "resume_api_http_request_duration_seconds_bucket" in response.text
    assert "resume_api_model_ready" in response.text
